"""Excel export: accepted identities, plus a sheet explaining every rejection."""

from __future__ import annotations

from collections.abc import Iterable
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font

from idcard_extractor.models import ARABIC_LABELS, IdentityRecord, ProcessedCard

REJECTED_HEADERS = ("رقم البطاقة", "نوع البطاقة", "ملف الصورة", "سبب الرفض")


def _setup_sheet(sheet, headers, width: int) -> None:
    sheet.sheet_view.rightToLeft = True
    sheet.append(list(headers))
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    sheet.freeze_panes = "A2"
    for column in sheet.iter_cols(min_row=1, max_row=1):
        sheet.column_dimensions[column[0].column_letter].width = width


def export_to_excel(
    records: Iterable[IdentityRecord],
    output_path: str | Path,
    rejected: Iterable[ProcessedCard] = (),
) -> Path:
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    workbook = Workbook()
    cards = workbook.active
    cards.title = "Cards"
    keys = list(ARABIC_LABELS)
    _setup_sheet(cards, [ARABIC_LABELS[k] for k in keys], width=18)

    for record in records:
        data = record.to_dict()
        # Identifiers stay strings so Excel does not turn them into numbers.
        cards.append([data.get(k) if data.get(k) is not None else "" for k in keys])
    for row in cards.iter_rows(min_row=2):
        for cell in row:
            if hasattr(cell.value, "year"):
                cell.number_format = "yyyy-mm-dd"

    sheet = workbook.create_sheet("Rejected")
    _setup_sheet(sheet, REJECTED_HEADERS, width=24)
    for card in rejected:
        sides = card.sides
        sheet.append([
            card.card_id,
            " + ".join(dict.fromkeys(s.card_class for s in sides)),
            " + ".join(dict.fromkeys(Path(s.source_image).name for s in sides)),
            card.reason,
        ])

    workbook.save(output_path)
    return output_path
