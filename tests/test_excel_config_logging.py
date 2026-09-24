from datetime import date

import pytest
from openpyxl import load_workbook

from idcard_extractor import logging_utils
from idcard_extractor.config import ConfigError, load_settings
from idcard_extractor.exporters.excel import export_to_excel
from idcard_extractor.models import IdentityRecord, ProcessedCard


def test_excel_export(tmp_path, make_front):
    record = IdentityRecord(first_name="أحمد", national_id="127901234567", date_of_birth=date(1979, 1, 5),
                            sex="ذكر", registration_number="A12345678")
    rejected = ProcessedCard("CARD_003", False, "لا يوجد ظهر", front=make_front("CARD_003"))
    path = export_to_excel([record], tmp_path / "out" / "cards.xlsx", [rejected])

    workbook = load_workbook(path)
    assert workbook.sheetnames == ["Cards", "Rejected"]
    cards = workbook["Cards"]
    headers = [cell.value for cell in cards[1]]
    assert headers[0] == "الاسم" and headers[-1] == "رقم القيد"
    values = dict(zip(headers, [cell.value for cell in cards[2]], strict=True))
    assert values["الرقم الوطني"] == "127901234567"  # kept as text
    assert values["تاريخ الميلاد"].date() == date(1979, 1, 5)
    assert values["رقم القيد"] == "A12345678"
    assert cards.sheet_view.rightToLeft

    row = [cell.value for cell in workbook["Rejected"][2]]
    assert row == ["CARD_003", "iraq id card", "CARD_003.jpg", "لا يوجد ظهر"]


def test_settings_from_env_file(tmp_path):
    env = tmp_path / "conf" / ".env"
    env.parent.mkdir()
    env.write_text("WEIGHTS_DIR=w\nDETECT_CONF=0.3\nYOLO_DEVICE=\nDB_FAILED_RECORDS_FILE=none\n", encoding="utf-8")
    settings = load_settings(env)
    assert settings.base_dir == env.parent.resolve()
    assert settings.findcard_weights == env.parent.resolve() / "w" / "findCard.pt"
    assert settings.detect_conf == 0.3
    assert settings.yolo_device is None
    assert settings.db_failed_records_file is None
    assert settings.db_mapping_file == env.parent.resolve() / "config" / "db_mapping.yaml"


def test_environment_overrides_env_file(tmp_path, monkeypatch):
    (tmp_path / ".env").write_text("OCR_DEVICE=gpu\n", encoding="utf-8")
    monkeypatch.setenv("OCR_DEVICE", "cpu")
    assert load_settings().ocr_device == "cpu"  # found from the working directory


def test_invalid_settings(tmp_path):
    with pytest.raises(ConfigError, match="not found"):
        load_settings(tmp_path / "missing.env")
    bad = tmp_path / "bad.env"
    bad.write_text("DETECT_CONF=high\n", encoding="utf-8")
    with pytest.raises(ConfigError, match="DETECT_CONF"):
        load_settings(bad)


def test_masking_helpers():
    assert logging_utils.mask("127901234567") == "12********67"
    assert logging_utils.mask_text("أحمد") == "أ***"
    message = "Duplicate entry '127901234567' for key 'national_no' (أحمد الكعبي)"
    assert logging_utils.redact(message) == "Duplicate entry '12********67' for key 'national_no' (*** ***)"

    logging_utils._mask_enabled = False
    assert logging_utils.redact(message) == message
